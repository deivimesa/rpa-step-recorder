"""Exportación a Word (.docx) con python-docx.

Genera el PDD con un diseño cuidado y legible para negocio + desarrollador RPA:
portada con resumen de fiabilidad, índice de pasos en tabla de 3 columnas y una
página por paso con jerarquía visual clara (cabecera + badge, descripción
destacada, capturas con pie de figura y datos agrupados en secciones).

No cambia la firma pública: export(session, out_path) -> Path.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from .. import doc_common as dc
from .. import storage


# ── constantes tipográficas y de color (sin números mágicos dispersos) ──
def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


ACCENT = _rgb(dc.ACCENT_HEX)
INK = _rgb(dc.INK_HEX)
MUTED = _rgb(dc.MUTED_HEX)
RULE = _rgb(dc.RULE_HEX)
REL_COLORS = {
    "alta": _rgb(dc.GREEN_HEX),
    "media": _rgb(dc.AMBER_HEX),
    "baja": _rgb(dc.RED_HEX),
}

# Tamaños de fuente (pt): se usan los tokens unificados de doc_common para que
# Word y PDF compartan la misma escala. FS_DESC se mapea al "lead" del documento.
FS_TITLE = dc.FS_TITLE
FS_SUBTITLE = dc.FS_SUBTITLE
FS_H1 = dc.FS_H1
FS_STEP = dc.FS_STEP
FS_DESC = dc.FS_LEAD
FS_SECTION = dc.FS_SECTION
FS_BODY = dc.FS_BODY
FS_SMALL = dc.FS_SMALL
FS_CAPTION = dc.FS_CAPTION

# Ancho útil de la página (pulgadas) para que las imágenes no se desborden.
CONTENT_W_IN = 6.3
DETAIL_W_IN = 3.6


def _set_cell_text(cell, text, bold=False, color=None, size=FS_BODY, align=None):
    """Escribe texto en una celda controlando negrita, color y tamaño."""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _shade_cell(cell, hex_fill: str):
    """Aplica un sombreado de fondo a una celda (XML, no hay API directa)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _accent_band(doc):
    """Dibuja una banda de acento fina bajo el título (párrafo con borde)."""
    band = doc.add_paragraph()
    band.paragraph_format.space_before = Pt(2)
    band.paragraph_format.space_after = Pt(10)
    p_pr = band._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24")  # grosor en octavos de punto (~3pt)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), dc.ACCENT_HEX)
    borders.append(bottom)
    p_pr.append(borders)
    return band


def _section_title(doc, text):
    """Subtítulo de sección con una regla fina debajo y espaciado uniforme."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = dc.LH_HEADING
    run = para.add_run(str(text).upper())
    run.bold = True
    run.font.name = dc.DOC_FONT_HEADING
    run.font.size = Pt(FS_SECTION)
    run.font.color.rgb = ACCENT
    # Regla fina (borde inferior gris claro) para separar la sección.
    p_pr = para._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), dc.RULE_HEX)
    borders.append(bottom)
    p_pr.append(borders)
    return para


def _kv_table(doc, rows):
    """Tabla clave/valor limpia: etiquetas en gris, valores en tinta."""
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.allow_autofit = False
    for key, value in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], key, bold=True, color=MUTED, size=FS_SMALL)
        _set_cell_text(cells[1], value, color=INK, size=FS_BODY)
        cells[0].width = Inches(1.9)
        cells[1].width = Inches(4.4)
    return table


def _caption(doc, text):
    """Pie de figura pequeño en gris, centrado bajo la imagen."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(1)
    para.paragraph_format.space_after = Pt(8)
    run = para.add_run(str(text))
    run.italic = True
    run.font.size = Pt(FS_CAPTION)
    run.font.color.rgb = MUTED


def _centered_picture(doc, path: Path, width_in: float):
    """Inserta una imagen centrada en su propio párrafo."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def _cover(doc, session):
    """Portada: título, banda de acento, tabla de datos y resumen de fiabilidad."""
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(60)
    title.paragraph_format.line_spacing = dc.LH_TITLE
    run = title.add_run(session.get("name", "Proceso"))
    run.bold = True
    run.font.name = dc.DOC_FONT_HEADING
    run.font.size = Pt(FS_TITLE)
    run.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(2)
    subtitle.paragraph_format.line_spacing = dc.LH_HEADING
    sub_run = subtitle.add_run(dc.DOC_SUBTITLE)
    sub_run.font.name = dc.DOC_FONT_HEADING
    sub_run.font.size = Pt(FS_SUBTITLE)
    sub_run.font.color.rgb = ACCENT

    _accent_band(doc)

    _kv_table(doc, dc.cover_rows(session))

    # Bloque "Resumen de fiabilidad" en la portada.
    _section_title(doc, "Resumen de fiabilidad")
    overview = doc.add_paragraph()
    overview.paragraph_format.line_spacing = dc.LH_BODY
    ov_run = overview.add_run(dc.reliability_overview(session))
    ov_run.bold = True
    ov_run.font.size = Pt(FS_BODY)
    ov_run.font.color.rgb = INK

    warning = dc.reliability_warning(session)
    if warning:
        warn_p = doc.add_paragraph()
        warn_run = warn_p.add_run(warning)
        warn_run.font.size = Pt(FS_SMALL)
        warn_run.font.color.rgb = REL_COLORS["baja"]


def _step_index(doc, session):
    """Índice/resumen de pasos como tabla de 3 columnas: # | Descripción | Fiabilidad."""
    heading = doc.add_paragraph()
    heading.paragraph_format.line_spacing = dc.LH_HEADING
    h_run = heading.add_run("Resumen de pasos")
    h_run.bold = True
    h_run.font.name = dc.DOC_FONT_HEADING
    h_run.font.size = Pt(FS_H1)
    h_run.font.color.rgb = INK

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light List Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False

    # Cabecera con fondo tenue y etiquetas claras.
    header = table.rows[0].cells
    for cell, label in zip(header, ("#", "Descripción", "Fiabilidad")):
        _set_cell_text(cell, label, bold=True, color=INK, size=FS_SMALL)
        _shade_cell(cell, "F2F2F2")
    header[0].width = Inches(0.5)
    header[1].width = Inches(4.6)
    header[2].width = Inches(1.2)

    for step in session.get("steps", []):
        cells = table.add_row().cells
        _set_cell_text(cells[0], str(step.get("index", "")), color=INK, size=FS_BODY)
        _set_cell_text(cells[1], dc.step_text(step), color=INK, size=FS_BODY)
        level = dc.step_level(step)
        label = dc.REL_SHORT.get(level, "—")
        _set_cell_text(
            cells[2], label, bold=True, color=REL_COLORS.get(level, MUTED), size=FS_BODY
        )
        cells[0].width = Inches(0.5)
        cells[1].width = Inches(4.6)
        cells[2].width = Inches(1.2)


def _step_header(doc, step):
    """Cabecera del paso: 'Paso N' en acento + badge de fiabilidad a la derecha."""
    table = doc.add_table(rows=1, cols=2)
    table.allow_autofit = False
    left, right = table.rows[0].cells
    _set_cell_text(left, f"Paso {step.get('index', '')}", bold=True, color=ACCENT, size=FS_STEP)
    # Encabezado del paso con la familia de títulos.
    for run in left.paragraphs[0].runs:
        run.font.name = dc.DOC_FONT_HEADING
    left.width = Inches(4.8)

    level = dc.step_level(step)
    if level:
        _set_cell_text(
            right,
            dc.REL_SHORT.get(level, level),
            bold=True,
            color=REL_COLORS.get(level, MUTED),
            size=FS_SMALL,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
    else:
        _set_cell_text(right, "", size=FS_SMALL)
    right.width = Inches(1.5)
    # Sin bordes en la cabecera (tabla solo para alinear a izquierda/derecha).


def _step_body(doc, session, step):
    """Contenido completo de un paso con jerarquía y secciones etiquetadas."""
    _step_header(doc, step)

    # Descripción destacada (tamaño mayor, interlineado de "lead").
    desc = doc.add_paragraph()
    desc.paragraph_format.space_before = Pt(4)
    desc.paragraph_format.line_spacing = dc.LH_LEAD
    desc_run = desc.add_run(dc.step_text(step))
    desc_run.font.size = Pt(FS_DESC)
    desc_run.font.color.rgb = INK

    # Fiabilidad detallada (motivo) debajo de la descripción.
    summary = dc.reliability_summary(step)
    if summary:
        level = dc.step_level(step)
        rel_p = doc.add_paragraph()
        rel_run = rel_p.add_run(f"Fiabilidad del selector: {summary}")
        rel_run.font.size = Pt(FS_SMALL)
        rel_run.bold = True
        rel_run.font.color.rgb = REL_COLORS.get(level, MUTED)

    # Capturas: pantalla completa y detalle del elemento, con pie de figura.
    full = step.get("full_image")
    if full:
        path = storage.image_path(session["id"], full)
        if path.exists():
            _centered_picture(doc, path, CONTENT_W_IN)
            _caption(doc, "Pantalla completa")

    zoom = step.get("zoom_image")
    if zoom:
        path = storage.image_path(session["id"], zoom)
        if path.exists():
            _centered_picture(doc, path, DETAIL_W_IN)
            _caption(doc, "Detalle del elemento")

    # Datos agrupados en secciones consistentes.
    ident = dc.identification_rows(step)
    if ident:
        _section_title(doc, "Identificación del elemento")
        _kv_table(doc, ident)

    location = dc.location_rows(step)
    if location:
        _section_title(doc, "Ubicación")
        _kv_table(doc, location)

    selectors = dc.selector_rows(step)
    if selectors:
        _section_title(doc, "Selectores para RPA")
        _kv_table(doc, selectors)


def _apply_footer(doc, session):
    """Nombre del proceso en gris pequeño en el pie de todas las páginas."""
    section = doc.sections[0]
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(session.get("name", "Proceso"))
    run.font.size = Pt(FS_CAPTION)
    run.font.color.rgb = MUTED


def _configure_base_styles(doc):
    """Fija la fuente de cuerpo (Calibri) y el tamaño base del estilo Normal."""
    normal = doc.styles["Normal"]
    normal.font.name = dc.DOC_FONT
    normal.font.size = Pt(dc.FS_BODY)


def export(session: dict, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    doc = Document()

    _configure_base_styles(doc)
    _apply_footer(doc, session)

    # ── portada ──
    _cover(doc, session)

    # ── índice / resumen de pasos ──
    doc.add_page_break()
    _step_index(doc, session)

    # ── una página por paso ──
    for step in session.get("steps", []):
        doc.add_page_break()
        _step_body(doc, session, step)

    doc.save(str(out_path))
    return out_path
